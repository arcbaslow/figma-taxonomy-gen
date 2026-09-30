"""Regression coverage for events shared by multiple Figma nodes."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest
from click.testing import CliRunner

from figma_taxonomy.cli import main
from figma_taxonomy.config import TaxonomyConfig
from figma_taxonomy.extractor import extract_elements
from figma_taxonomy.mcp_tools import (
    export_taxonomy_tool,
    extract_taxonomy_tool,
    validate_taxonomy_tool,
)
from figma_taxonomy.models import ScreenElement, TaxonomyEvent
from figma_taxonomy.output.json_schema import write_json
from figma_taxonomy.taxonomy_engine import generate_taxonomy
from figma_taxonomy.validate import _events_from_dict, diff_taxonomy_dicts


@pytest.fixture
def variant_file() -> Path:
    return Path(__file__).parent / "fixtures" / "variant_sources.json"


def test_variants_keep_shared_and_unique_controls(variant_file: Path) -> None:
    config = TaxonomyConfig()
    elements = extract_elements(json.loads(variant_file.read_text()), config)
    assert {element.node_id for element in elements} == {"1:1", "1:2", "2:1", "2:2"}
    events = {event.event_name: event for event in generate_taxonomy(elements, config)}
    assert events["settings_save_clicked"].source_node_ids == ["1:1", "1:2", "2:1"]
    assert events["settings_save_clicked"].source_node_id == "1:1"
    assert events["settings_cancel_clicked"].source_node_ids == ["2:2"]


@pytest.mark.parametrize("fmt", ["json", "csv", "markdown", "excel"])
def test_mcp_export_preserves_every_source(
    variant_file: Path, tmp_path: Path, fmt: str,
) -> None:
    result = extract_taxonomy_tool(str(variant_file))
    event = next(e for e in result["events"] if e["event_name"] == "settings_save_clicked")
    assert event["source_node_ids"] == ["1:1", "1:2", "2:1"]
    path = tmp_path / ("out.xlsx" if fmt == "excel" else f"out.{fmt}")
    export_taxonomy_tool(result, fmt, str(path))
    if fmt == "json":
        body = json.loads(path.read_text())["events"][event["event_name"]]
        assert body["source"] == "figma:node_id:1:1"
        assert body["sources"] == [f"figma:node_id:{node}" for node in event["source_node_ids"]]
        assert _events_from_dict({event["event_name"]: body})[0].source_node_ids == event["source_node_ids"]
    elif fmt == "csv":
        with path.open(encoding="utf-8", newline="") as stream:
            rows = [row for row in csv.DictReader(stream) if row["Event Type"] == event["event_name"]]
        assert rows
        assert all(row["Source Node ID"] == "1:1" for row in rows)
        assert all(json.loads(row["Source Node IDs"]) == event["source_node_ids"] for row in rows)
    elif fmt == "excel":
        workbook = openpyxl.load_workbook(path)
        sheet = workbook["Events"]
        assert sheet.cell(2, 15).value == "Source Node IDs"
        assert sheet.cell(3, 14).value == "1:1"
        assert json.loads(sheet.cell(3, 15).value) == event["source_node_ids"]
        workbook.close()
    else:
        content = path.read_text(encoding="utf-8")
        for node in event["source_node_ids"]:
            assert f"`{node}`" in content


def test_model_accepts_legacy_and_multiple_sources() -> None:
    legacy = TaxonomyEvent("click", "", "", source_node_id="1:1")
    assert legacy.source_node_ids == ["1:1"]
    event = TaxonomyEvent("click", "", "", source_node_ids=["2:1", "2:1", "", "3:1"])
    assert event.source_node_id == "2:1"
    assert event.source_node_ids == ["2:1", "3:1"]


@pytest.mark.parametrize("style", ["snake_case", "camelCase"])
def test_truncation_collision_is_actionable(style: str) -> None:
    config = TaxonomyConfig()
    config.naming.style = style
    config.naming.max_event_length = 12
    elements = [
        ScreenElement(node, "settings", label, "button", None, True)
        for node, label in [("1:1", "save_address"), ("1:2", "save_profile")]
    ]
    with pytest.raises(ValueError, match="max_event_length") as error:
        generate_taxonomy(elements, config)
    assert "1:1" in str(error.value) and "1:2" in str(error.value)


def _body(*ids: str) -> dict:
    return {"source": f"figma:node_id:{ids[0]}" if ids else "",
            "sources": [f"figma:node_id:{node}" for node in ids], "properties": {}}


def test_secondary_source_matches_rename_and_reports_removal() -> None:
    report = diff_taxonomy_dicts({"old": _body("1:1", "2:1")}, {"new": _body("2:1")})
    assert report.renamed == [("old", "new")]
    assert not report.added and not report.removed
    assert report.source_changes == [{"event_name": "new", "added": [], "removed": ["1:1"]}]


def test_source_order_is_not_drift() -> None:
    assert diff_taxonomy_dicts(
        {"save": _body("1:1", "2:1")}, {"save": _body("2:1", "1:1")},
    ).is_clean()


def test_source_addition_and_replacement_are_drift() -> None:
    report = diff_taxonomy_dicts({"save": _body("1:1", "2:1")}, {"save": _body("1:1", "3:1")})
    assert not report.is_clean()
    assert report.source_changes == [{"event_name": "save", "added": ["3:1"], "removed": ["2:1"]}]


@pytest.mark.parametrize("reverse", [False, True])
def test_ambiguous_split_or_merge_is_not_a_rename(reverse: bool) -> None:
    old = {"combined": _body("1:1", "2:1")}
    new = {"first": _body("1:1"), "second": _body("2:1")}
    if reverse:
        old, new = new, old
    report = diff_taxonomy_dicts(old, new)
    assert not report.renamed
    assert {event.event_name for event in report.added} == set(new)
    assert set(report.removed) == set(old)


def test_legacy_source_loads_without_migration_drift() -> None:
    old = {"save": {"source": "figma:node_id:1:1", "properties": {}}}
    assert diff_taxonomy_dicts(old, {"save": _body("1:1")}).is_clean()


def test_json_writer_round_trip(variant_file: Path, tmp_path: Path) -> None:
    config = TaxonomyConfig()
    events = generate_taxonomy(extract_elements(json.loads(variant_file.read_text()), config), config)
    output = tmp_path / "taxonomy.json"
    write_json(events, config, output)
    stored = json.loads(output.read_text())
    hydrated = _events_from_dict(stored["events"])
    assert [e.source_node_ids for e in hydrated] == [e.source_node_ids for e in events]
    assert validate_taxonomy_tool(stored, str(variant_file))["is_clean"]
    stored["events"]["settings_save_clicked"]["sources"].remove("figma:node_id:1:2")
    report = validate_taxonomy_tool(stored, str(variant_file))
    assert not report["is_clean"]
    assert report["source_changes"] == [
        {"event_name": "settings_save_clicked", "added": ["1:2"], "removed": []},
    ]


@pytest.mark.parametrize("command", ["validate", "diff"])
def test_cli_reports_source_only_drift(
    variant_file: Path, tmp_path: Path, command: str,
) -> None:
    current_path = tmp_path / "current.json"
    export_taxonomy_tool(extract_taxonomy_tool(str(variant_file)), "json", str(current_path))
    stored = json.loads(current_path.read_text())
    stored["events"]["settings_save_clicked"]["sources"].remove("figma:node_id:1:2")
    old_path = tmp_path / "old.json"
    old_path.write_text(json.dumps(stored), encoding="utf-8")
    target = [str(current_path)] if command == "diff" else ["--fixture", str(variant_file)]
    result = CliRunner().invoke(main, [command, str(old_path), *target, "--exit-code"])
    assert result.exit_code == 1
    assert "Source changes (1)" in result.output
    assert "+ node 1:2" in result.output


def test_cli_collision_fails_before_writing_outputs(variant_file: Path, tmp_path: Path) -> None:
    config_path = tmp_path / "naming.yaml"
    config_path.write_text("naming:\n  max_event_length: 8\n", encoding="utf-8")
    output = tmp_path / "output"
    result = CliRunner().invoke(main, [
        "extract", "--fixture", str(variant_file), "--config", str(config_path),
        "--output", str(output),
    ])
    assert result.exit_code == 1
    assert "Error: Event names" in result.output
    assert "naming.max_event_length" in result.output
    assert not output.exists()


def test_enrichment_keeps_grouped_sources(variant_file: Path) -> None:
    from figma_taxonomy.ai_enricher import enrich_events

    config = TaxonomyConfig()
    events = generate_taxonomy(extract_elements(json.loads(variant_file.read_text()), config), config)
    client = Mock()
    client.messages.create.return_value = SimpleNamespace(content=[SimpleNamespace(text=json.dumps({
        "suggestions": [{"event_name": "settings_save_clicked", "properties": [
            {"name": "save_mode", "type": "string", "enum": ["manual", "auto"]},
        ]}],
    }))])
    enriched = enrich_events(events, config, client)
    client.messages.create.assert_called_once()
    prompt = client.messages.create.call_args.kwargs["messages"][0]["content"]
    assert prompt.count("- settings_save_clicked:") == 1
    save = next(e for e in enriched if e.event_name == "settings_save_clicked")
    assert save.source_node_ids == ["1:1", "1:2", "2:1"]
    assert any(p.name == "save_mode" for p in save.properties)


def test_pageview_truncation_collision_is_rejected() -> None:
    config = TaxonomyConfig()
    config.naming.pattern = "{screen}_{action}"
    config.naming.max_event_length = 4
    with pytest.raises(ValueError, match="truncate"):
        generate_taxonomy([ScreenElement("1:1", "home", "save", "button", None, True)], config)


def test_custom_pattern_merges_sources_without_duplicate_ids() -> None:
    config = TaxonomyConfig()
    config.naming.pattern = "{element}_{action}"
    elements = [
        ScreenElement("1:1", "home", "save", "button", None, True),
        ScreenElement("2:1", "settings", "save", "button", None, True),
        ScreenElement("1:1", "home", "save", "button", None, True),
    ]
    events = generate_taxonomy(elements, config)
    assert [e.event_name for e in events] == ["save_clicked", "pageview"]
    assert events[0].source_node_ids == ["1:1", "2:1"]
