import json
from pathlib import Path

import openpyxl
import pytest

from figma_taxonomy.config import TaxonomyConfig
from figma_taxonomy.models import EventProperty, TaxonomyEvent
from figma_taxonomy.output.excel import write_excel


def test_excel_preserves_all_schemas_sources_and_literal_text(tmp_path: Path) -> None:
    event = TaxonomyEvent("=1+1", "=SUM(1,2)", "Счёт", [
        EventProperty(f"p{i}", "string", "=HYPERLINK(\"https://example.com\")", ["да", "нет"])
        for i in range(7)
    ], source_node_ids=["1:1", "2:2"])
    path = tmp_path / "taxonomy.xlsx"
    write_excel([event], TaxonomyConfig(), path)
    book = openpyxl.load_workbook(path)
    try:
        assert book["Events"].cell(3, 3).value == "=1+1"
        assert book["Events"].cell(3, 3).data_type == "s"
        rows = list(book["Event Properties"].values)
        assert len(rows) == 8
        assert rows[-1][2:4] == ("p6", "string")
        assert json.loads(rows[-1][5]) == ["да", "нет"]
        assert json.loads(rows[-1][6]) == ["1:1", "2:2"]
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    finally:
        book.close()


@pytest.mark.parametrize("text", ["x" * 32768, "invalid\x00text"], ids=["too-long", "control-character"])
def test_unrepresentable_excel_text_fails_before_replacing_file(tmp_path: Path, text: str) -> None:
    path = tmp_path / "taxonomy.xlsx"
    path.write_bytes(b"original")
    with pytest.raises(ValueError, match="Excel"):
        write_excel([TaxonomyEvent("event", "Flow", text)], TaxonomyConfig(), path)
    assert path.read_bytes() == b"original"
